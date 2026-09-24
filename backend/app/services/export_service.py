"""
export_service.py -- exports Excel/PDF du rapport Direction (Par ligne, F5 Écritures
Odoo) et, depuis le 23/09, de l'historique des scans opérateur.

*** REFONTE 2026-09-23 *** : délègue la mise en page à export_helpers.py (pied de page
Vusine + date/heure + pagination, arrondis, en-tête Excel figé avec filtres) -- les
signatures publiques ne changent pas, reports_routes.py et labo_routes.py n'ont rien à
modifier.
"""
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from . import export_helpers as h

COLONNES = ["Ligne", "Réel", "Théorique", "Performance", "Palettes", "Arrêt (min)"]


def _df_rapport(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame([{
        "Ligne": f"{r['code']} — {r['nom']}",
        "Réel": r["reel_total"],
        "Théorique": r["theorique_total"],
        "Performance": (r["performance_moyenne"] / 100) if r["performance_moyenne"] is not None else None,
        "Palettes": r["nb_palettes"],
        "Arrêt (min)": r["temps_arret_min"],
    } for r in rows], columns=COLONNES)


def generer_excel_rapport(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    df = _df_rapport(rows)
    contenu = h.build_excel("Rapport par ligne", df, sous_titre=f"Du {date_debut.isoformat()} au {date_fin.isoformat()}")
    # Performance en vrai pourcentage Excel (0.95 affiché "95 %"), pas un texte "95%".
    wb = load_workbook(BytesIO(contenu))
    ws = wb.active
    col_perf = COLONNES.index("Performance") + 1
    for ligne in range(6, ws.max_row + 1):
        cell = ws.cell(row=ligne, column=col_perf)
        if cell.value is not None:
            cell.number_format = "0 %"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generer_pdf_rapport(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    df = _df_rapport(rows)
    df["Performance"] = df["Performance"].apply(lambda v: None if pd.isna(v) else round(v * 100))
    df = df.rename(columns={"Performance": "Performance (%)"})
    return h.build_pdf("Rapport par ligne", df, sous_titre=f"Période : {date_debut.isoformat()} au {date_fin.isoformat()}",
                        decimales={"Performance (%)": 0})


def generer_csv_rapport(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    df = _df_rapport(rows)
    df["Performance"] = df["Performance"].apply(lambda v: None if pd.isna(v) else round(v * 100, 1))
    df = df.rename(columns={"Performance": "Performance (%)"})
    return h.build_csv(df)


# =============================================================
# *** AJOUT (chantier Labo, F5 export) *** : ce que Vusine écrirait dans Odoo (palettes
# réelles) si F5 était actif, et sa comparaison avec ce qu'Odoo contient réellement
# (of_cache) -- F5 lui-même n'écrit jamais dans Odoo (décision actée), ces exports sont
# la matérialisation de "l'aperçu" en attendant que la fiabilité soit prouvée.
# =============================================================

COLONNES_ECRITURES = ["Ligne", "Jour", "Produit", "Qté réelle (palettes)", "Nb palettes", "Dont partielles"]
COLONNES_COMPARAISON = ["Ligne", "Jour", "Produit", "Qté Vusine", "Qté Odoo", "Écart", "Écart %"]


def _df_ecritures(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame([{
        "Ligne": r["ligne_code"], "Jour": r["jour"].isoformat(), "Produit": r["produit_nom"],
        "Qté réelle (palettes)": r["quantite_reelle"], "Nb palettes": r["nb_palettes"],
        "Dont partielles": r["nb_palettes_partielles"],
    } for r in rows], columns=COLONNES_ECRITURES)


def _df_comparaison(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame([{
        "Ligne": r["ligne_code"], "Jour": r["jour"].isoformat(), "Produit": r["produit_nom"],
        "Qté Vusine": r["quantite_vusine"], "Qté Odoo": r["quantite_odoo"],
        "Écart": r["ecart_qte"], "Écart %": r["ecart_pct"],
    } for r in rows], columns=COLONNES_COMPARAISON)


def generer_excel_ecritures_proposees(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_excel("Écritures proposées vers Odoo", _df_ecritures(rows),
                          sous_titre=f"F5 non actif -- du {date_debut.isoformat()} au {date_fin.isoformat()}")


def generer_excel_comparaison_odoo(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_excel("Comparaison Vusine (palettes) vs Odoo", _df_comparaison(rows),
                          sous_titre=f"Du {date_debut.isoformat()} au {date_fin.isoformat()}",
                          decimales={"Écart %": 1})


def generer_pdf_ecritures_proposees(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_pdf("Écritures proposées vers Odoo (F5 non actif)", _df_ecritures(rows),
                        sous_titre=f"Période : {date_debut.isoformat()} au {date_fin.isoformat()}")


def generer_pdf_comparaison_odoo(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_pdf("Comparaison Vusine (palettes) vs Odoo", _df_comparaison(rows),
                        sous_titre=f"Période : {date_debut.isoformat()} au {date_fin.isoformat()}",
                        decimales={"Écart %": 1})


def generer_csv_ecritures_proposees(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_csv(_df_ecritures(rows))


def generer_csv_comparaison_odoo(rows: list[dict], date_debut: date, date_fin: date) -> bytes:
    return h.build_csv(_df_comparaison(rows))


# =============================================================
# *** AJOUT 2026-09-23 *** : historique des scans (palettes) par opérateur -- écran
# Rapports (Direction, tous opérateurs) et écran tablette "Mon historique" (opérateur,
# ses propres scans). cf. reports_service.historique_scans pour la restriction d'accès.
# =============================================================

COLONNES_HISTORIQUE = ["Date / heure", "Opérateur", "Matricule", "Ligne", "Produit",
                        "N° lot", "Cartons", "Colisage", "Quantité", "Rebuts", "Statut"]


def _df_historique_scans(rows: list) -> pd.DataFrame:
    return pd.DataFrame([{
        "Date / heure": r.created_at.strftime("%d/%m/%Y %H:%M"),
        "Opérateur": r.operateur_nom, "Matricule": r.operateur_matricule or "—",
        "Ligne": r.ligne_code, "Produit": r.produit_nom or "—", "N° lot": r.numero_lot,
        "Cartons": r.nb_cartons, "Colisage": r.colisage_carton, "Quantité": r.quantite_totale,
        "Rebuts": r.nb_rebuts,
        "Statut": "Complète" if r.complete else f"Partielle ({r.motif_partielle or 'motif non précisé'})",
    } for r in rows], columns=COLONNES_HISTORIQUE)


def generer_excel_historique_scans(rows: list, date_debut: date, date_fin: date) -> bytes:
    return h.build_excel("Historique des scans", _df_historique_scans(rows),
                          sous_titre=f"Du {date_debut.isoformat()} au {date_fin.isoformat()}")


def generer_pdf_historique_scans(rows: list, date_debut: date, date_fin: date) -> bytes:
    return h.build_pdf("Historique des scans", _df_historique_scans(rows),
                        sous_titre=f"Période : {date_debut.isoformat()} au {date_fin.isoformat()}")


def generer_csv_historique_scans(rows: list, date_debut: date, date_fin: date) -> bytes:
    return h.build_csv(_df_historique_scans(rows))


# =============================================================
# *** AJOUT 2026-09-24 (Palier 0) *** : Pareto des causes d'arrêt.
# `pareto` est un ParetoArretsOut (cf. pertes_service.pareto_arrets).
# =============================================================

COLONNES_PARETO = ["Rang", "Cause", "Durée (min)", "Nb arrêts", "% du total", "% cumulé", "Coût estimé (FCFA)"]
DECIMALES_PARETO = {"% du total": 1, "% cumulé": 1}


def _df_pareto(pareto) -> pd.DataFrame:
    return pd.DataFrame([{
        "Rang": c.rang, "Cause": c.cause, "Durée (min)": c.duree_min, "Nb arrêts": c.nb_arrets,
        "% du total": c.pct, "% cumulé": c.pct_cumule,
        "Coût estimé (FCFA)": c.cout_fcfa,
    } for c in pareto.causes], columns=COLONNES_PARETO).astype({"Coût estimé (FCFA)": "Int64"})
    # Int64 « nullable » : un montant entier, vide (et non « 0 » ni « 3000,0 ») quand le
    # coût n'est pas calculable.


def _sous_titre_pareto(pareto) -> str:
    total = f"Total : {pareto.total_duree_min} min sur {pareto.total_nb_arrets} arrêt(s)"
    if pareto.total_cout_fcfa is not None:
        total += f" — coût estimé {pareto.total_cout_fcfa:,} FCFA".replace(",", " ")
    periode = f"Du {pareto.date_debut.isoformat()} au {pareto.date_fin.isoformat()}"
    note = (f"Coût = production planifiée perdue pendant les heures de poste, valorisée au « "
            f"{pareto.libelle_valeur} » ; vide si le planning ou la valeur manque.")
    return f"{periode}. {total}. {note}"


def generer_excel_pareto(pareto) -> bytes:
    return h.build_excel("Pareto des arrêts", _df_pareto(pareto),
                         sous_titre=_sous_titre_pareto(pareto), decimales=DECIMALES_PARETO)


def generer_pdf_pareto(pareto) -> bytes:
    return h.build_pdf("Pareto des arrêts", _df_pareto(pareto),
                       sous_titre=_sous_titre_pareto(pareto), decimales=DECIMALES_PARETO)


def generer_csv_pareto(pareto) -> bytes:
    return h.build_csv(_df_pareto(pareto))


# =============================================================
# *** AJOUT 2026-09-24 (Palier 1) *** : TRS décomposé. `trs` est un TrsOut
# (cf. trs_service.calculer_trs) ; la ligne de total « USINE » clôt le tableau.
# =============================================================

COLONNES_TRS = ["Ligne", "Jours", "Planifié", "Conforme", "Rebuts", "Arrêt (min)",
                "Disponibilité %", "Performance %", "Qualité %", "TRS %",
                "Pertes arrêts (pcs)", "Pertes cadence (pcs)", "Pertes rebuts (pcs)"]
DECIMALES_TRS = {"Disponibilité %": 1, "Performance %": 1, "Qualité %": 1, "TRS %": 1}


def _df_trs(trs) -> pd.DataFrame:
    def ligne(l):
        return {
            "Ligne": f"{l.code} — {l.nom}" if l.ligne_id else "USINE (total)",
            "Jours": l.jours, "Planifié": l.qte_planifiee, "Conforme": l.production_conforme,
            "Rebuts": l.rebuts, "Arrêt (min)": l.minutes_arret,
            "Disponibilité %": l.disponibilite_pct, "Performance %": l.performance_pct,
            "Qualité %": l.qualite_pct, "TRS %": l.trs_pct,
            "Pertes arrêts (pcs)": l.pertes_arrets_pieces, "Pertes cadence (pcs)": l.pertes_cadence_pieces,
            "Pertes rebuts (pcs)": l.pertes_rebuts_pieces,
        }
    df = pd.DataFrame([ligne(l) for l in trs.lignes] + [ligne(trs.usine)], columns=COLONNES_TRS)
    for col in ("Disponibilité %", "Performance %", "Qualité %", "TRS %"):
        df[col] = df[col].astype(float)
    return df


def _sous_titre_trs(trs) -> str:
    texte = (f"Du {trs.date_debut.isoformat()} au {trs.date_fin.isoformat()} — {trs.nb_jours} jour(s) complet(s). "
             f"TRS = Disponibilité x Performance x Qualité = production conforme / quantité planifiée. "
             f"Cible : {str(trs.cible_pct).rstrip('0').rstrip('.')} %.")
    if not trs.qualite_renseignee:
        texte += " Qualité : aucun rebut déclaré sur la période -> 100 % par défaut, non mesurée."
    if trs.pieces_hors_planning:
        texte += f" {trs.pieces_hors_planning} pièces produites hors planning, non comptées."
    return texte


def generer_excel_trs(trs) -> bytes:
    return h.build_excel("TRS", _df_trs(trs), sous_titre=_sous_titre_trs(trs), decimales=DECIMALES_TRS)


def generer_pdf_trs(trs) -> bytes:
    return h.build_pdf("TRS (Taux de Rendement Synthétique)", _df_trs(trs), sous_titre=_sous_titre_trs(trs), decimales=DECIMALES_TRS)


def generer_csv_trs(trs) -> bytes:
    return h.build_csv(_df_trs(trs))


# =============================================================
# *** AJOUT 2026-09-24 (Palier 1) *** : changements de série (SMED). `smed` = SmedOut.
# =============================================================

COLONNES_SMED = ["Jour", "Ligne", "Produit avant", "Produit après", "Dernier scan avant",
                 "Premier scan après", "Écart entre scans (min)", "Changement déclaré (min)", "Objectif"]


def _df_smed(smed) -> pd.DataFrame:
    return pd.DataFrame([{
        "Jour": c.jour.strftime("%d/%m/%Y"), "Ligne": c.ligne_code,
        "Produit avant": c.produit_avant, "Produit après": c.produit_apres,
        "Dernier scan avant": c.dernier_scan_avant.strftime("%H:%M"),
        "Premier scan après": c.premier_scan_apres.strftime("%H:%M"),
        "Écart entre scans (min)": c.ecart_scans_min, "Changement déclaré (min)": c.arret_declare_min,
        "Objectif": "Dépassé" if c.au_dessus_objectif else ("OK" if smed.objectif_min else "—"),
    } for c in smed.changements], columns=COLONNES_SMED).astype({"Changement déclaré (min)": "Int64"})


def _sous_titre_smed(smed) -> str:
    texte = f"Du {smed.date_debut.isoformat()} au {smed.date_fin.isoformat()} — {smed.nb_changements} changement(s)."
    if smed.moyenne_min is not None:
        texte += f" Moyenne {smed.moyenne_min} min, médiane {smed.mediane_min} min, meilleur {smed.meilleur_min} min."
    if smed.objectif_min:
        texte += f" Objectif : {smed.objectif_min} min."
    return texte + (" L'écart entre scans est une borne haute : il inclut le remplissage de la 1re palette du produit suivant,"
                    " hors pause et hors heures de poste.")


def generer_excel_smed(smed) -> bytes:
    return h.build_excel("Changements de série", _df_smed(smed), sous_titre=_sous_titre_smed(smed))


def generer_pdf_smed(smed) -> bytes:
    return h.build_pdf("Changements de série", _df_smed(smed), sous_titre=_sous_titre_smed(smed))


def generer_csv_smed(smed) -> bytes:
    return h.build_csv(_df_smed(smed))
